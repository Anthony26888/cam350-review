"""Build print-ready A4 User Guide: docs/USER_GUIDE.md -> exports/.

Outputs:
  exports/USER_GUIDE_A4.docx
  exports/USER_GUIDE_A4.pdf

Usage:
  python tools/build_user_guide.py
"""
import os
import re
import sys

SRC = os.path.join("docs", "USER_GUIDE.md")
OUT_DOCX = os.path.join("exports", "USER_GUIDE_A4.docx")
OUT_PDF = os.path.join("exports", "USER_GUIDE_A4.pdf")
FOOTER_TEXT = "CAM350 Review Assistant v2.2.3 \u2014 User Guide"


# ---------------------------------------------------------------- parsing
def parse_markdown(text):
    """Yield blocks: ('h', level, text) | ('p', text) | ('ul', [items])
    | ('ol', [items]) | ('code', lang, text) | ('table', header, rows)
    | ('quote', text) | ('hr',) | ('img', alt)."""
    lines = text.splitlines()
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            lang = s[3:].strip()
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1  # closing fence
            yield ("code", lang, "\n".join(buf))
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            yield ("h", len(m.group(1)), m.group(2).strip())
            i += 1
            continue
        if s.startswith("|") and i + 1 < n and re.match(
            r"^\|?[\s\-:|]+\|?\s*$", lines[i + 1].strip()
        ):
            header = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(
                    [c.strip() for c in lines[i].strip().strip("|").split("|")]
                )
                i += 1
            yield ("table", header, rows)
            continue
        if re.match(r"^(\d+)[.)]\s+", s):
            items = []
            while i < n and re.match(r"^(\d+)[.)]\s+", lines[i].strip()):
                items.append(re.sub(r"^(\d+)[.)]\s+", "", lines[i].strip()))
                i += 1
            yield ("ol", items)
            continue
        if re.match(r"^[-*]\s+", s):
            items = []
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            yield ("ul", items)
            continue
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            yield ("quote", " ".join(buf))
            continue
        if re.match(r"^---+\s*$", s):
            yield ("hr",)
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            yield ("img", m.group(1) or m.group(2))
            i += 1
            continue
        # paragraph: join wrapped lines
        buf = [s]
        i += 1
        while i < n and lines[i].strip() and not re.match(
            r"^(#{1,4}\s|\||\d+[.)]\s|[-*]\s|>|```|---+|!\[)", lines[i].strip()
        ):
            buf.append(lines[i].strip())
            i += 1
        yield ("p", " ".join(buf))


def strip_md(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
    return t


def md_runs_docx(para, text):
    """Add runs with **bold** and `code` handling."""
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = run.font.size
        else:
            para.add_run(part)


def md_to_reportlab(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(
        r"`(.+?)`", r'<font face="Courier">\1</font>', t
    )
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
    return t


# ---------------------------------------------------------------- DOCX
def build_docx(blocks):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(25)
        section.right_margin = Mm(20)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for hname, size, color in (
        ("Heading 1", 18, "1E293B"),
        ("Heading 2", 14, "1E293B"),
        ("Heading 3", 12, "334155"),
    ):
        hs = doc.styles[hname]
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = RGBColor.from_string(color)

    # footer with page number field
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = FOOTER_TEXT + "\t"
    fld1 = OxmlElement("w:fldSimple")
    fld1.set(qn("w:instr"), "PAGE")
    footer._p.append(fld1)
    footer.paragraph_format.tab_stops.add_tab_stop(Mm(170))

    for b in blocks:
        kind = b[0]
        if kind == "h":
            _, level, text = b
            if level == 1 and text.startswith("CAM350") is False and "User Guide" in text:
                p = doc.add_paragraph()
                run = p.add_run(strip_md(text))
                run.bold = True
                run.font.size = Pt(22)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(strip_md(text), level=min(level, 3))
        elif kind == "p":
            p = doc.add_paragraph()
            md_runs_docx(p, b[1])
        elif kind in ("ul", "ol"):
            for item in b[1]:
                style_name = "List Bullet" if kind == "ul" else "List Number"
                p = doc.add_paragraph(style=style_name)
                md_runs_docx(p, item)
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b[2])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif kind == "quote":
            p = doc.add_paragraph()
            run = p.add_run(strip_md(b[1]))
            run.italic = True
        elif kind == "img":
            p = doc.add_paragraph()
            run = p.add_run("[Figure: " + b[1] + "]")
            run.italic = True
        elif kind == "hr":
            doc.add_paragraph("─" * 40)
        elif kind == "table":
            _, header, rows = b
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, hcell in enumerate(header):
                cell = table.cell(0, j)
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(strip_md(hcell))
                run.bold = True
            for ri, row in enumerate(rows, start=1):
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    table.cell(ri, j).text = strip_md(val)
    return doc


# ---------------------------------------------------------------- PDF
def build_pdf(blocks):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        Preformatted, HRFlowable, ListFlowable, ListItem, PageBreak,
    )
    from reportlab.lib.enums import TA_CENTER

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Times-Roman", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(25 * mm, 12 * mm, FOOTER_TEXT)
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, f"{doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        OUT_PDF, pagesize=A4,
        topMargin=20 * mm, bottomMargin=20 * mm,
        leftMargin=25 * mm, rightMargin=20 * mm,
        title="CAM350 Review Assistant — User Guide",
        author="CAM350 Review",
    )
    styles = {
        "title": ParagraphStyle("title", fontName="Times-Bold", fontSize=22,
                                alignment=TA_CENTER, spaceAfter=6),
        "h1": ParagraphStyle("h1", fontName="Times-Bold", fontSize=17,
                             textColor=colors.HexColor("#1E293B"),
                             spaceBefore=14, spaceAfter=6),
        "h2": ParagraphStyle("h2", fontName="Times-Bold", fontSize=13,
                             textColor=colors.HexColor("#1E293B"),
                             spaceBefore=10, spaceAfter=4),
        "h3": ParagraphStyle("h3", fontName="Times-Bold", fontSize=11.5,
                             textColor=colors.HexColor("#334155"),
                             spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("body", fontName="Times-Roman", fontSize=10.5,
                               leading=14.5, spaceAfter=5),
        "bullet": ParagraphStyle("bullet", parent=None, fontName="Times-Roman",
                                 fontSize=10.5, leading=14.5, leftIndent=14,
                                 spaceAfter=3),
        "code": ParagraphStyle("code", fontName="Courier", fontSize=8.5,
                               leading=11.5, leftIndent=8, spaceAfter=6),
        "quote": ParagraphStyle("quote", fontName="Times-Italic", fontSize=10.5,
                                leading=14, leftIndent=10,
                                textColor=colors.HexColor("#475569"),
                                spaceAfter=6),
        "cell": ParagraphStyle("cell", fontName="Times-Roman", fontSize=9.5,
                               leading=12),
        "cellH": ParagraphStyle("cellH", fontName="Times-Bold", fontSize=9.5,
                                leading=12),
    }
    usable = A4[0] - 45 * mm
    story = []
    for b in blocks:
        kind = b[0]
        if kind == "h":
            _, level, text = b
            t = strip_md(text)
            if level == 1:
                story.append(Paragraph(t.replace("&", "&amp;"), styles["title"]))
            else:
                story.append(Paragraph(md_to_reportlab(text),
                                       styles[f"h{min(level, 3)}"]))
        elif kind == "p":
            story.append(Paragraph(md_to_reportlab(b[1]), styles["body"]))
        elif kind in ("ul", "ol"):
            items = [ListItem(Paragraph(md_to_reportlab(x), styles["bullet"]),
                              bulletColor=colors.black,
                              value="•" if kind == "ul" else n + 1)
                     for n, x in enumerate(b[1])]
            story.append(ListFlowable(items, bulletType="bullet" if kind == "ul" else "1",
                                      leftIndent=14))
            story.append(Spacer(1, 4))
        elif kind == "code":
            story.append(Preformatted(b[2], styles["code"]))
        elif kind == "quote":
            story.append(Paragraph(md_to_reportlab(b[1]), styles["quote"]))
        elif kind == "img":
            story.append(Paragraph(f"<i>[Figure: {b[1]}]</i>", styles["quote"]))
        elif kind == "hr":
            story.append(HRFlowable(width="40%", thickness=0.5,
                                    color=colors.HexColor("#94A3B8")))
            story.append(Spacer(1, 6))
        elif kind == "table":
            _, header, rows = b
            ncol = len(header)
            widths = None
            if ncol == 2:
                widths = [usable * 0.32, usable * 0.68]
            elif ncol == 3:
                widths = [usable * 0.20, usable * 0.24, usable * 0.56]
            data = [[Paragraph(f"<b>{strip_md(c)}</b>", styles["cellH"])
                     for c in header]]
            for row in rows:
                cells = [(row[j] if j < len(row) else "") for j in range(ncol)]
                data.append([Paragraph(md_to_reportlab(strip_md(c)), styles["cell"])
                             for c in cells])
            t = Table(data, colWidths=widths, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#F8FAFC")]),
            ]))
            story.append(t)
            story.append(Spacer(1, 6))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def main():
    with open(SRC, encoding="utf-8") as f:
        text = f.read()
    blocks = list(parse_markdown(text))
    os.makedirs("exports", exist_ok=True)
    doc = build_docx(blocks)
    doc.save(OUT_DOCX)
    print(f"wrote {OUT_DOCX}")
    build_pdf(blocks)
    print(f"wrote {OUT_PDF}")


if __name__ == "__main__":
    sys.exit(main())
