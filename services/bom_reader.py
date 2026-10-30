import re
from typing import Dict, Any, List, Optional, Tuple, Callable

import openpyxl

from models.bom import BomRow


BOM_FIELDS = ("designator", "mpn", "description", "quantity", "note")

BOM_REQUIRED = ["designator"]

BOM_FIELD_ALIASES: Dict[str, tuple] = {
    "designator": (
        "designator", "refdes", "ref des", "ref", "reference",
        "component", "part reference", "reference designator",
        "ref des list", "designators", "refs",
    ),
    "mpn": (
        "mpn", "part number", "partnumber", "part no", "part no.",
        "pn", "part", "manufacturer part number", "mpn no", "mfr part no",
    ),
    "description": (
        "description", "desc", "mô tả", "mo ta", "comment", "value",
    ),
    "quantity": (
        "quantity", "qty", "q'ty", "q ty", "sl", "số lượng", "so luong",
        "amount", "count",
    ),
    "note": (
        "note", "notes", "ghi chú", "ghi chu", "remark", "remarks",
    ),
}

_RANGE_LIMIT = 5000
_RANGE_RE = re.compile(r"^([A-Za-z_]*?)(\d+)\s*-\s*([A-Za-z_]*?)(\d+)$")
_SPLIT_RE = re.compile(r"[,;+/\|&\n]+")
_JUNK_RE = re.compile(r"[\(\)\[\]\{\}\"\']")
_QTY_SUFFIX_RE = re.compile(r"\s*[x×\*]\s*\d+\s*$")


def _normalize(value: str) -> str:
    return "".join(value.lower().split())


def guess_bom_mapping(headers: List[str]) -> Dict[str, str]:
    normalized = {_normalize(h): h for h in headers if h}
    mapping: Dict[str, str] = {}
    for field in BOM_FIELDS:
        for alias in BOM_FIELD_ALIASES[field]:
            key = _normalize(alias)
            if key in normalized:
                mapping[field] = normalized[key]
                break
    return mapping


def clean_designators(raw: Any) -> Tuple[str, List[str]]:
    """Clean a raw BOM Designator cell.

    Returns (display, expanded) where display joins expanded items with a
    single comma. Expands same-prefix ranges (C1-C4 -> C1..C4), unifies
    separators to commas, strips junk marks, keeps zero padding.
    """
    if raw is None:
        return "", []
    text = str(raw).strip()
    if not text:
        return "", []
    text = re.sub(r"[\s]+", " ", text)
    text = _SPLIT_RE.sub(",", text)
    tokens: List[str] = []
    for tok in text.split(","):
        tok = tok.strip().strip(".:")
        if not tok:
            continue
        tok = re.sub(r"\([^)]*\)", "", tok).strip()
        tok = _JUNK_RE.sub("", tok).strip()
        tok = _QTY_SUFFIX_RE.sub("", tok).strip()
        tok = re.sub(r"\s+", "", tok)
        if tok:
            tokens.append(tok)
    expanded: List[str] = []
    for tok in tokens:
        expanded.extend(_expand_token(tok))
    # de-dup preserving order (case-insensitive)
    seen = set()
    unique: List[str] = []
    for d in expanded:
        key = d.upper()
        if key not in seen:
            seen.add(key)
            unique.append(d)
    return ",".join(unique), unique


def _expand_token(tok: str) -> List[str]:
    m = _RANGE_RE.match(tok)
    if not m:
        return [tok]
    pre_a, num_a, pre_b, num_b = m.group(1), m.group(2), m.group(3), m.group(4)
    if pre_a.upper() != pre_b.upper():
        return [tok]
    a, b = int(num_a), int(num_b)
    if a > b:
        a, b = b, a
    if b - a > _RANGE_LIMIT:
        return [tok]
    width = max(len(num_a), len(num_b))
    prefix = pre_a or pre_b
    # keep original case of the first non-empty prefix
    return [f"{prefix}{str(i).zfill(width)}" for i in range(a, b + 1)]


def _safe_int(value: Any) -> Tuple[int, str]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0, ""
    raw = str(value).strip()
    try:
        return int(float(raw)), raw
    except (ValueError, TypeError):
        return 0, raw


class BomReader:

    @staticmethod
    def read_headers(file_path: str) -> List[str]:
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise FileNotFoundError(f"File not found: {file_path}")
        except Exception as e:
            raise ValueError(f"Cannot open Excel file: {e}")
        try:
            sheet = workbook.active
            headers: List[str] = []
            if sheet is not None:
                for row in sheet.iter_rows(max_row=1, values_only=True):
                    headers = [str(h).strip() if h is not None else "" for h in row]
                    break
        finally:
            workbook.close()
        return headers

    @staticmethod
    def read_with_mapping(
        file_path: str,
        mapping: Optional[Dict[str, str]] = None,
        progress: Optional[Callable[[int, int], None]] = None,
    ) -> List[BomRow]:
        mapping = mapping or {}
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise FileNotFoundError(f"File not found: {file_path}")
        except Exception as e:
            raise ValueError(f"Cannot open Excel file: {e}")
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("Excel file has no active sheet")

            row_iter = sheet.iter_rows(values_only=True)
            try:
                header_row = next(row_iter)
            except StopIteration:
                raise ValueError("Excel file has no data rows")

            headers = [str(h).strip() if h is not None else "" for h in header_row]
            header_lookup = {h.lower(): i for i, h in enumerate(headers) if h}

            col_map: Dict[str, int] = {}
            for field in BOM_FIELDS:
                target = mapping.get(field)
                if not target:
                    continue
                key = str(target).strip().lower()
                if key in header_lookup:
                    col_map[field] = header_lookup[key]
            if "designator" not in col_map:
                raise ValueError("BOM mapping needs a Designator column")

            result: List[BomRow] = []
            stt = 0
            total = max(int(getattr(sheet, "max_row", 0) or 0) - 1, 0)
            row_count = 0
            for row_idx, row in enumerate(row_iter, start=1):
                row_count = row_idx
                if progress is not None and (row_idx % 200 == 0 or row_idx == total):
                    progress(row_idx, total)

                if all(c is None or str(c).strip() == "" for c in row):
                    continue
                raw_des = _cell_str(row, col_map.get("designator"))
                if not raw_des:
                    continue
                stt += 1
                display, expanded = clean_designators(raw_des)
                qty, qty_raw = _safe_int(_cell_raw(row, col_map.get("quantity")))
                result.append(BomRow(
                    stt=stt,
                    designator_raw=raw_des,
                    designator_display=display,
                    designators=expanded,
                    mpn=_cell_str(row, col_map.get("mpn")),
                    description=_cell_str(row, col_map.get("description")),
                    quantity=qty,
                    quantity_raw=qty_raw,
                    note=_cell_str(row, col_map.get("note")),
                ))

            if row_count == 0:
                raise ValueError("Excel file has no data rows")
            if progress is not None and row_count != total:
                final = max(total, row_count)
                progress(final, final)
            return result
        finally:
            workbook.close()


def _cell_raw(row: tuple, col_idx: Optional[int]) -> Any:
    if col_idx is None or col_idx >= len(row):
        return None
    return row[col_idx]


def _cell_str(row: tuple, col_idx: Optional[int]) -> str:
    val = _cell_raw(row, col_idx)
    if val is None:
        return ""
    return str(val).strip()
