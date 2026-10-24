import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import types

from PySide6.QtWidgets import QApplication

from models.review import ReviewRecord
from services.eco_diff import EcoAdded, EcoRemoved, EcoMoved, EcoMpnChanged, EcoDiff
from ui.eco_diff_dialog import EcoResultDialog
from ui.main_window import MainWindow


def _app():
    return QApplication.instance() or QApplication([])


class _Repo:
    def __init__(self):
        self._next_id = 100

    def insert(self, record):
        self._next_id += 1
        return self._next_id

    def update(self, record):
        pass


class _HistoryRepo:
    def __init__(self):
        self.entries = []

    def add(self, entry):
        self.entries.append(entry)
        return len(self.entries)

    def list_by(self, *args, **kwargs):
        return []


class _Table:
    def set_records(self, records):
        pass

    def select_record(self, index):
        pass


class _Panel:
    def set_record_list(self, records):
        pass

    def display_record(self, *args):
        pass

    def reload_history(self, entries):
        pass

    def update_stats(self, *args):
        pass


def _holder(records):
    _app()
    h = types.SimpleNamespace()
    h._records = records
    h._current_index = 0
    h._dirty = False
    h._repo = _Repo()
    h._history_repo = _HistoryRepo()
    h._table_widget = _Table()
    h._review_panel = _Panel()
    h._jump_popup = None
    h._gerber_viewer = None
    h._status_label = types.SimpleNamespace(text="",
        setText=lambda t: setattr(h._status_label, "text", t))
    h._undo_stack = []
    h._redo_stack = []
    h._undo_action = None
    h._redo_action = None
    for name in ["_apply_eco_diff", "_push_undo", "_log_history",
                 "_record_snapshot", "_update_undo_actions",
                 "_sync_after_mutation", "_select_and_display",
                 "_refresh_history_panel", "_update_progress",
                 "_on_eco_locate"]:
        setattr(h, name, getattr(MainWindow, name).__get__(h))
    return h


def _records():
    return [
        ReviewRecord(designator="C1", mpn="M-A", layer="Top",
                     old_x=1, old_y=2, old_rotation=0, status="OK",
                     new_x=1, new_y=2, new_rotation=0),
        ReviewRecord(designator="C2", mpn="M-B", layer="Top",
                     old_x=3, old_y=4, old_rotation=0, status="Pending"),
    ]


def _diff():
    return EcoDiff(
        added=[EcoAdded(designator="C9", mpn="M-N", layer="Top",
                        x=9, y=9, rotation=0)],
        removed=[EcoRemoved(designator="C2", mpn="M-B", layer="Top",
                            x=3, y=4, rotation=0)],
        moved=[],
        mpn_changed=[EcoMpnChanged(designator="C1", old_mpn="M-A",
                                   new_mpn="M-A2")],
    )


def test_apply_eco_diff_all_groups():
    h = _holder(_records())
    changed = h._apply_eco_diff(_diff())
    assert changed == 3
    by_des = {r.designator: r for r in h._records}
    # added as Pending
    assert by_des["C9"].status == "Pending"
    assert (by_des["C9"].old_x, by_des["C9"].old_y) == (9, 9)
    # removed kept + remarked
    assert "ECO removed" in by_des["C2"].remark
    # mpn updated
    assert by_des["C1"].mpn == "M-A2"
    actions = sorted(e.action for e in h._history_repo.entries)
    assert actions == ["ECO", "ECO", "Import"]
    assert len(h._undo_stack) == 1  # single undo step


def test_apply_eco_moved_resets_review():
    recs = [ReviewRecord(designator="C1", mpn="M", layer="Top",
                          old_x=1, old_y=2, old_rotation=0, status="Edited",
                          new_x=1.5, new_y=2, new_rotation=0)]
    h = _holder(recs)
    diff = EcoDiff(moved=[EcoMoved(
        designator="C1", old_x=1.5, old_y=2, old_rotation=0,
        new_x=10, new_y=20, new_rotation=90, dx=8.5, dy=18, drot=90)])
    h._apply_eco_diff(diff)
    r = h._records[0]
    assert (r.old_x, r.old_y, r.old_rotation) == (10, 20, 90)
    assert r.new_x is None and r.status == "Pending"
    assert h._history_repo.entries[0].action == "Edit"


def test_result_dialog_groups_and_locate():
    _app()
    dlg = EcoResultDialog(_diff(), 2, 2)
    assert dlg._table.rowCount() == 1  # first non-empty group: added
    dlg._show_group("mpn")
    assert dlg._table.rowCount() == 1
    assert dlg._table.item(0, 0).text() == "C1"
    got = []
    dlg.locate_requested.connect(got.append)
    dlg._show_group("removed")
    dlg._on_cell_clicked(0, 0)
    assert got == ["C2"]
    dlg._show_group("added")
    dlg._on_cell_clicked(0, 0)
    assert got == ["C2"]  # added rows do not locate (not in session yet)
    dlg.close()


def test_result_dialog_duplicates_group():
    _app()
    from services.eco_diff import EcoDuplicate
    diff = _diff()
    diff.new_duplicates.append(EcoDuplicate(designator="C1", side="new", count=3))
    diff.old_unique, diff.new_unique = 2, 2
    dlg = EcoResultDialog(diff, 2, 3)
    dlg._show_group("dup")
    assert dlg._table.rowCount() == 1
    assert dlg._table.item(0, 0).text() == "C1"
    assert dlg._table.item(0, 1).text() == "New"
    assert dlg._table.item(0, 2).text() == "×3"
    dlg.close()


def test_export_eco_report_includes_duplicates(tmp_path):
    from services.export_service import ExportService
    import openpyxl
    from services.eco_diff import EcoDuplicate
    diff = _diff()
    diff.new_duplicates.append(EcoDuplicate(designator="C1", side="new", count=2))
    path = tmp_path / "eco.xlsx"
    ExportService.export_eco_report(diff, str(path))
    wb = openpyxl.load_workbook(str(path))
    ws = wb.active
    groups = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert "Duplicates" in groups
    assert ws.cell(row=ws.max_row, column=2).value == "C1"
    wb.close()


def test_on_eco_locate_selects_record():
    h = _holder(_records())
    h._on_eco_locate("c1")
    assert h._current_index == 0
    h._on_eco_locate("NOPE")
    assert h._current_index == 0
